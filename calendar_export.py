#!/usr/bin/env python3
"""calendar_export.py - turn a contentwright batch .md into an xlsx content calendar.

Usage:
    python calendar_export.py batch.md calendar.xlsx

Expects the batch heading format produced by contentwright:
    ## 3. LinkedIn - Wed 16 Apr - contrarian take

Anything it cannot parse is still carried into the sheet with blank slot fields,
so nothing is silently dropped.

Requires openpyxl.
"""

import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl missing. Install with: pip install openpyxl --break-system-packages", file=sys.stderr)
    sys.exit(2)

HEADING = re.compile(r"^##\s+(?:(\d+)\.\s*)?(.+)$")
HOOKS = re.compile(r"\*\*Hook options:?\*\*\s*(.*)", re.I)
NOTES = re.compile(r"\*\*Notes:?\*\*\s*(.*)", re.I)

PRIMARY = "5DADE2"
SECONDARY = "708090"
DIVIDER = "D3D3D3"


def parse(md: str):
    pieces = []
    current = None
    for line in md.split("\n"):
        h = HEADING.match(line)
        if h:
            if current:
                pieces.append(current)
            title = h.group(2).strip()
            parts = [p.strip() for p in re.split(r"\s+-\s+", title)]
            current = {
                "n": h.group(1) or "",
                "platform": parts[0] if parts else title,
                "slot": parts[1] if len(parts) > 1 else "",
                "format": parts[2] if len(parts) > 2 else "",
                "hook": "",
                "notes": "",
                "body_words": 0,
            }
            continue
        if current is None:
            continue
        hk = HOOKS.search(line)
        if hk and not current["hook"]:
            current["hook"] = hk.group(1).strip()
            continue
        nt = NOTES.search(line)
        if nt:
            current["notes"] = nt.group(1).strip()
            continue
        if not line.strip().startswith("**"):
            current["body_words"] += len(re.findall(r"\b[\w'-]+\b", line))
    if current:
        pieces.append(current)
    return pieces


def write_xlsx(pieces, out_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Content calendar"

    headers = ["#", "Slot", "Platform", "Format", "Hook", "Words", "Status", "Notes"]
    widths = [5, 16, 14, 22, 52, 8, 12, 46]

    head_fill = PatternFill("solid", fgColor=PRIMARY)
    head_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11)
    note_font = Font(name="Calibri", size=11, color=SECONDARY)
    thin = Side(style="thin", color=DIVIDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for i, p in enumerate(pieces, start=2):
        row = [p["n"], p["slot"], p["platform"], p["format"], p["hook"],
               p["body_words"], "Draft", p["notes"]]
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.font = note_font if col == 8 else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=col in (5, 8))
            cell.border = border

    if pieces:
        ws.auto_filter.ref = f"A1:H{len(pieces) + 1}"
    wb.save(out_path)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        return 2
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.is_file():
        print(f"file not found: {src}", file=sys.stderr)
        return 2
    pieces = parse(src.read_text(encoding="utf-8"))
    if not pieces:
        print("no '## n. Platform - slot - format' headings found. Nothing to export.", file=sys.stderr)
        return 1
    write_xlsx(pieces, dst)
    print(f"wrote {dst} with {len(pieces)} pieces")
    return 0


if __name__ == "__main__":
    sys.exit(main())
